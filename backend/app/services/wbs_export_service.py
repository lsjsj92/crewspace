# backend/app/services/wbs_export_service.py
# Epic 기준 WBS Excel 내보내기 서비스

from __future__ import annotations

import asyncio
import io
import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_app_config
from app.models.card import Card, CardLink, CardType
from app.models.user import User
from app.repositories.card_repository import CardRepository
from app.repositories.project_repository import ProjectRepository
from app.services.project_permission_service import check_project_permission
from app.utils.datetime_utils import today_kst

# 간트 영역 좌측 정보 컬럼 수 (구분/번호/업무명/담당자/시작일/종료일/상태)
INFO_COLUMNS = ["구분", "번호", "업무명", "담당자", "시작일", "종료일", "상태"]
GANTT_START_COL = len(INFO_COLUMNS) + 1

# 카드 타입별 간트 바 색상 (frontend constants/cardTypes.ts와 동일 팔레트, ARGB)
BAR_COLORS = {
    "epic": "FF722ED1",
    "story": "FF1890FF",
    "task": "FF52C41A",
    "sub_task": "FFFAAD14",
}
COMPLETED_BAR_COLOR = "FFBFBFBF"
WEEKEND_FILL_COLOR = "FFF5F5F5"
HEADER_FILL_COLOR = "FFFAFAFA"
DEFAULT_BAR_COLOR = "FF8C8C8C"

TYPE_LABELS = {"epic": "Epic", "story": "Story", "task": "Task", "sub_task": "Sub-task"}


@dataclass
class WbsRow:
    type_label: str
    number: str
    title: str
    assignees: str
    start: date | None
    due: date | None
    status: str
    indent: int
    card_type: str
    is_done: bool


@dataclass
class WbsSheetData:
    title: str
    epic_number: str
    epic_title: str
    epic_start: date | None
    epic_due: date | None
    rows: list[WbsRow] = field(default_factory=list)


def _card_type_value(card: Card) -> str:
    return card.card_type.value if isinstance(card.card_type, CardType) else card.card_type


def _card_status(card: Card) -> str:
    """카드 상태 문자열: 취소/완료가 우선, 그 외에는 현재 컬럼명."""
    if card.cancelled_at is not None:
        return "취소"
    if card.completed_at is not None:
        return "완료"
    return card.column.name if card.column else ""


def _assignee_names(card: Card) -> str:
    names = []
    for a in card.assignees or []:
        if a.user:
            names.append(a.user.display_name or a.user.username or "")
    return ", ".join(n for n in names if n)


def _to_row(card: Card, prefix: str, indent: int) -> WbsRow:
    card_type = _card_type_value(card)
    return WbsRow(
        type_label=TYPE_LABELS.get(card_type, card_type),
        number=f"{prefix}-{card.card_number}" if prefix else str(card.card_number),
        title=card.title,
        assignees=_assignee_names(card),
        start=card.start_date,
        due=card.due_date,
        status=_card_status(card),
        indent=indent,
        card_type=card_type,
        is_done=card.completed_at is not None or card.cancelled_at is not None,
    )


def _start_sort_key(card: Card) -> tuple[date, int]:
    """시작일 오름차순, 시작일 없는 카드는 마지막, 동률은 카드 번호순."""
    return (card.start_date or date.max, card.card_number)


def _collect_epic_rows(epic: Card, children_map: dict[UUID, list[Card]], prefix: str) -> list[WbsRow]:
    """Epic 직계 자식(Story/Task)을 시작일순으로 배치하고 Story 하위 Task를 들여쓰기로 포함한다."""
    rows: list[WbsRow] = []
    for child in sorted(children_map.get(epic.id, []), key=_start_sort_key):
        rows.append(_to_row(child, prefix, indent=0))
        if _card_type_value(child) == "story":
            for task in sorted(children_map.get(child.id, []), key=_start_sort_key):
                rows.append(_to_row(task, prefix, indent=1))
    return rows


def _sanitize_sheet_title(title: str, used: set[str]) -> str:
    """Excel 시트명 제약(31자, 특수문자 금지)에 맞게 정리하고 중복을 방지한다."""
    cleaned = re.sub(r"[\\/*?:\[\]]", " ", title).strip() or "Epic"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 2
    while candidate in used:
        tail = f"~{suffix}"
        candidate = cleaned[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate)
    return candidate


def _epic_sheet(
    epic: Card,
    children_map: dict[UUID, list[Card]],
    prefix: str,
    used_titles: set[str],
) -> WbsSheetData:
    number = f"{prefix}-{epic.card_number}" if prefix else str(epic.card_number)
    return WbsSheetData(
        title=_sanitize_sheet_title(f"{number} {epic.title}", used_titles),
        epic_number=number,
        epic_title=epic.title,
        epic_start=epic.start_date,
        epic_due=epic.due_date,
        rows=_collect_epic_rows(epic, children_map, prefix),
    )


def _orphan_sheet(
    cards: list[Card], linked_card_ids: set[UUID], prefix: str, used_titles: set[str]
) -> WbsSheetData | None:
    """Epic에 속하지 않고 보조 연결도 없는 독립 카드(Story/Task)를 '기타' 시트로 모은다."""
    orphans = sorted(
        (
            c for c in cards
            if c.parent_id is None
            and c.id not in linked_card_ids
            and _card_type_value(c) not in ("epic", "sub_task")
        ),
        key=_start_sort_key,
    )
    if not orphans:
        return None
    return WbsSheetData(
        title=_sanitize_sheet_title("기타 (Epic 없음)", used_titles),
        epic_number="",
        epic_title="기타 (Epic 없음)",
        epic_start=None,
        epic_due=None,
        rows=[_to_row(c, prefix, indent=0) for c in orphans],
    )


def _build_sheet_data(
    cards: list[Card], prefix: str, links: list[CardLink]
) -> list[WbsSheetData]:
    """카드 목록을 Epic 단위 시트 데이터로 변환한다.

    보조 연결(다중 부모)된 카드는 연결된 모든 상위(Epic/Story) 아래에도 표기된다.
    """
    children_map: dict[UUID, list[Card]] = {}
    for card in cards:
        if card.parent_id is not None:
            children_map.setdefault(card.parent_id, []).append(card)

    # 보조 연결된 카드를 연결 상위의 자식 목록에도 추가한다 (같은 상위 내 중복 방지)
    cards_by_id = {c.id: c for c in cards}
    for link in links:
        child = cards_by_id.get(link.card_id)
        if child is None:
            continue
        bucket = children_map.setdefault(link.parent_id, [])
        if all(c.id != child.id for c in bucket):
            bucket.append(child)

    used_titles: set[str] = set()
    epics = sorted((c for c in cards if _card_type_value(c) == "epic"), key=_start_sort_key)
    sheets = [_epic_sheet(e, children_map, prefix, used_titles) for e in epics]
    orphan = _orphan_sheet(cards, {link.card_id for link in links}, prefix, used_titles)
    if orphan:
        sheets.append(orphan)
    return sheets


def _sheet_dates(sheet: WbsSheetData, max_days: int) -> list[date]:
    """시트의 간트 날짜 범위를 계산한다. Epic과 모든 행의 날짜를 포괄한다."""
    all_dates = [d for d in (sheet.epic_start, sheet.epic_due) if d is not None]
    for row in sheet.rows:
        all_dates.extend(d for d in (row.start, row.due) if d is not None)
    if not all_dates:
        start = today_kst()
        return [start + timedelta(days=i) for i in range(30)]
    start, end = min(all_dates), max(all_dates)
    total_days = min((end - start).days + 1, max_days)
    return [start + timedelta(days=i) for i in range(max(total_days, 1))]


def _write_epic_title_cell(ws: Worksheet, sheet: WbsSheetData) -> None:
    period = ""
    if sheet.epic_start or sheet.epic_due:
        period = f" ({sheet.epic_start or ''} ~ {sheet.epic_due or ''})"
    epic_label = f"Epic: {sheet.epic_number} {sheet.epic_title}{period}".strip()
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(INFO_COLUMNS))
    ws.cell(row=1, column=1, value=epic_label).font = Font(bold=True, size=12)


def _write_month_header(ws: Worksheet, dates: list[date]) -> None:
    """1행 간트 영역: 같은 연월의 일자 열을 병합해 월 라벨을 표시한다."""
    header_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
    col = GANTT_START_COL
    while col - GANTT_START_COL < len(dates):
        d = dates[col - GANTT_START_COL]
        span = sum(1 for x in dates[col - GANTT_START_COL:] if (x.year, x.month) == (d.year, d.month))
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + span - 1)
        cell = ws.cell(row=1, column=col, value=f"{d.year}-{d.month:02d}")
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = header_fill
        col += span


def _write_column_headers(ws: Worksheet, dates: list[date]) -> None:
    """2행: 정보 컬럼 제목 + 일자 헤더 (주말은 회색 배경)."""
    header_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
    weekend_fill = PatternFill(start_color=WEEKEND_FILL_COLOR, end_color=WEEKEND_FILL_COLOR, fill_type="solid")
    for idx, name in enumerate(INFO_COLUMNS, start=1):
        cell = ws.cell(row=2, column=idx, value=name)
        cell.font = Font(bold=True, size=9)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for idx, d in enumerate(dates):
        cell = ws.cell(row=2, column=GANTT_START_COL + idx, value=d.day)
        cell.font = Font(size=8)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = weekend_fill if d.weekday() >= 5 else header_fill


def _write_sheet_header(ws: Worksheet, sheet: WbsSheetData, dates: list[date]) -> None:
    _write_epic_title_cell(ws, sheet)
    _write_month_header(ws, dates)
    _write_column_headers(ws, dates)


def _write_data_row(ws: Worksheet, row_idx: int, row: WbsRow, dates: list[date]) -> None:
    weekend_fill = PatternFill(start_color=WEEKEND_FILL_COLOR, end_color=WEEKEND_FILL_COLOR, fill_type="solid")
    bar_color = COMPLETED_BAR_COLOR if row.is_done else BAR_COLORS.get(row.card_type, DEFAULT_BAR_COLOR)
    bar_fill = PatternFill(start_color=bar_color, end_color=bar_color, fill_type="solid")

    values = [
        row.type_label,
        row.number,
        ("    " * row.indent) + row.title,
        row.assignees,
        row.start.isoformat() if row.start else "",
        row.due.isoformat() if row.due else "",
        row.status,
    ]
    for idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=idx, value=value)
        cell.font = Font(size=9, bold=(row.indent == 0 and row.card_type == "story"))

    bar_start = row.start or row.due
    bar_end = row.due or row.start
    for idx, d in enumerate(dates):
        cell = ws.cell(row=row_idx, column=GANTT_START_COL + idx)
        if bar_start is not None and bar_end is not None and bar_start <= d <= bar_end:
            cell.fill = bar_fill
        elif d.weekday() >= 5:
            cell.fill = weekend_fill


def _apply_column_widths(ws: Worksheet, date_count: int) -> None:
    info_widths = [8, 10, 40, 14, 11, 11, 8]
    for idx, width in enumerate(info_widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for idx in range(date_count):
        ws.column_dimensions[get_column_letter(GANTT_START_COL + idx)].width = 3

    # 정보 컬럼과 헤더 행 고정
    ws.freeze_panes = ws.cell(row=3, column=GANTT_START_COL)


def _build_workbook(sheets: list[WbsSheetData], max_days: int) -> bytes:
    """시트 데이터로 Excel 워크북을 생성한다 (동기 함수, to_thread로 호출)."""
    wb = Workbook()
    wb.remove(wb.active)

    if not sheets:
        ws = wb.create_sheet("WBS")
        ws.cell(row=1, column=1, value="내보낼 Epic 카드가 없습니다.")
    for sheet in sheets:
        ws = wb.create_sheet(sheet.title)
        dates = _sheet_dates(sheet, max_days)
        _write_sheet_header(ws, sheet, dates)
        for offset, row in enumerate(sheet.rows):
            _write_data_row(ws, 3 + offset, row, dates)
        if not sheet.rows:
            ws.cell(row=3, column=1, value="하위 카드가 없습니다.").font = Font(size=9, italic=True)
        _apply_column_widths(ws, len(dates))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def export_project_wbs(
    db: AsyncSession,
    project_id: UUID,
    current_user: User,
) -> tuple[bytes, str]:
    """프로젝트 카드를 Epic별 시트의 WBS Excel로 생성해 (파일 내용, 파일명)을 반환한다."""
    await check_project_permission(db, project_id, current_user, ["manager", "member", "viewer"])

    project = await ProjectRepository(db).get_by_id(project_id)
    prefix = project.prefix if project else ""
    project_name = project.name if project else "project"

    repo = CardRepository(db)
    cards = await repo.get_project_cards(project_id, include_archived=True)
    links = await repo.get_project_card_links(project_id)

    app_config = get_app_config()
    sheets = _build_sheet_data(cards, prefix, links)
    content = await asyncio.to_thread(_build_workbook, sheets, app_config.wbs_max_gantt_days)

    filename = f"{prefix or project_name}_WBS_{today_kst().isoformat()}.xlsx"
    return content, filename
